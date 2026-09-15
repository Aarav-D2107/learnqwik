"""
Document parsing — turns an uploaded file into a list of pages/slides.

Never "PDF -> text -> LLM". Each page is inspected on its own and classified so
the router downstream can decide whether Vision is actually needed:

    text_rich   plenty of selectable text, no meaningful visuals  -> text model only
    sparse      very little text for the page area                -> vision
    scanned     effectively no text but the page has image content-> vision
    mixed       real text AND meaningful visuals                  -> BOTH
    empty       nothing at all                                    -> skipped

Optional dependencies are imported lazily so a missing wheel degrades one file
type rather than crashing the whole service.
"""
import io
import logging
import os

import config

log = logging.getLogger("learnqwik.documents")

QUALITY_TEXT_RICH = "text_rich"
QUALITY_SPARSE = "sparse"
QUALITY_SCANNED = "scanned"
QUALITY_MIXED = "mixed"
QUALITY_EMPTY = "empty"

# A page with fewer characters than this is suspicious regardless of visuals.
SPARSE_CHAR_THRESHOLD = 180
SCANNED_CHAR_THRESHOLD = 40


class UnsupportedFileError(ValueError):
    pass


class DocumentParseError(RuntimeError):
    pass


def detect_extension(filename, declared_mime=None):
    ext = (os.path.splitext(filename or "")[1] or "").lstrip(".").lower()
    if ext not in config.ALLOWED_EXTENSIONS:
        raise UnsupportedFileError(
            ".%s is not a supported file type. Supported: %s"
            % (ext or "?", ", ".join(sorted(config.ALLOWED_EXTENSIONS)))
        )
    return ext


def is_source_file(ext):
    return ext in config.SOURCE_EXTENSIONS


def classify_page(text, has_visuals, visual_area_ratio=0.0):
    chars = len((text or "").strip())
    if chars == 0 and not has_visuals:
        return QUALITY_EMPTY
    if chars < SCANNED_CHAR_THRESHOLD and has_visuals:
        return QUALITY_SCANNED
    if chars < SPARSE_CHAR_THRESHOLD:
        return QUALITY_SPARSE
    if has_visuals and visual_area_ratio >= 0.12:
        return QUALITY_MIXED
    return QUALITY_TEXT_RICH


def needs_vision(quality):
    return quality in (QUALITY_SPARSE, QUALITY_SCANNED, QUALITY_MIXED)


# ------------------------------------------------------------------ PDF
def _parse_pdf(data):
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise DocumentParseError(
            "PDF support requires PyMuPDF. Install backend/requirements.txt."
        )
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception:
        raise DocumentParseError("This PDF could not be opened. It may be corrupt or password protected.")

    pages = []
    try:
        for index, page in enumerate(doc, start=1):
            text = page.get_text("text") or ""
            try:
                images = page.get_images(full=True)
            except Exception:
                images = []
            drawings = 0
            try:
                drawings = len(page.get_drawings())
            except Exception:
                pass

            page_area = max(1.0, page.rect.width * page.rect.height)
            visual_area = 0.0
            for img in images:
                try:
                    for rect in page.get_image_rects(img[0]):
                        visual_area += rect.width * rect.height
                except Exception:
                    visual_area += page_area * 0.15
            ratio = min(1.0, visual_area / page_area)
            has_visuals = bool(images) or drawings > 12

            pages.append({
                "page_number": index,
                "extracted_text": text.strip(),
                "has_visuals": has_visuals,
                "visual_area_ratio": round(ratio, 3),
                "content_quality": classify_page(text, has_visuals, ratio),
                "_render": ("pdf", index - 1),
            })
    finally:
        doc.close()
    return pages


def render_pdf_page(data, page_index, dpi=None):
    """Rasterize one page to PNG bytes for the Vision model."""
    import fitz
    dpi = dpi or config.VISION_RENDER_DPI
    doc = fitz.open(stream=data, filetype="pdf")
    try:
        page = doc.load_page(page_index)
        pix = page.get_pixmap(matrix=fitz.Matrix(dpi / 72.0, dpi / 72.0), alpha=False)
        return pix.tobytes("png")
    finally:
        doc.close()


# ------------------------------------------------------------------ PPTX
def _parse_pptx(data):
    try:
        from pptx import Presentation
    except ImportError:
        raise DocumentParseError("PPTX support requires python-pptx.")
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception:
        raise DocumentParseError("This presentation could not be opened.")

    pages = []
    for index, slide in enumerate(prs.slides, start=1):
        texts = []
        picture_count = 0
        table_count = 0
        chart_count = 0
        for shape in slide.shapes:
            try:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    texts.append(shape.text_frame.text.strip())
                if shape.shape_type is not None and "PICTURE" in str(shape.shape_type):
                    picture_count += 1
                if getattr(shape, "has_table", False):
                    table_count += 1
                    for row in shape.table.rows:
                        texts.append(" | ".join(c.text.strip() for c in row.cells))
                if getattr(shape, "has_chart", False):
                    chart_count += 1
            except Exception:
                continue
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                texts.append("[Speaker notes] " + slide.notes_slide.notes_text_frame.text.strip())
        except Exception:
            pass

        text = "\n".join(texts)
        has_visuals = bool(picture_count or chart_count or table_count)
        ratio = 0.4 if (picture_count or chart_count) else (0.2 if table_count else 0.0)
        pages.append({
            "page_number": index,
            "extracted_text": text,
            "has_visuals": has_visuals,
            "visual_area_ratio": ratio,
            "content_quality": classify_page(text, has_visuals, ratio),
            "_render": ("pptx-embedded", index - 1),
        })
    return pages


def extract_pptx_images(data, slide_index):
    """Slides can't be rasterized without a renderer, so we send the embedded
    images themselves to the Vision model instead."""
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    slides = list(prs.slides)
    if slide_index >= len(slides):
        return []
    out = []
    for shape in slides[slide_index].shapes:
        try:
            if shape.shape_type is not None and "PICTURE" in str(shape.shape_type):
                image = shape.image
                out.append((image.blob, "image/%s" % (image.ext or "png")))
        except Exception:
            continue
    return out[:3]


# ------------------------------------------------------------------ DOCX
def _parse_docx(data):
    try:
        import docx
    except ImportError:
        raise DocumentParseError("DOCX support requires python-docx.")
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception:
        raise DocumentParseError(
            "This document could not be opened. Legacy .doc files must be saved as .docx first."
        )

    blocks = []
    for para in document.paragraphs:
        if para.text.strip():
            style = (para.style.name or "").lower() if para.style else ""
            prefix = "## " if "heading" in style else ""
            blocks.append(prefix + para.text.strip())
    for table in document.tables:
        blocks.append("[Table]")
        for row in table.rows:
            blocks.append(" | ".join(c.text.strip() for c in row.cells))

    image_count = 0
    try:
        image_count = sum(1 for r in document.part.rels.values() if "image" in r.reltype)
    except Exception:
        pass

    # Word has no fixed pages; chunk into pseudo-pages of ~2500 characters so
    # page references stay meaningful to the student.
    text = "\n".join(blocks)
    pages = []
    size = 2500
    chunks = [text[i:i + size] for i in range(0, max(len(text), 1), size)] or [""]
    for index, chunk in enumerate(chunks, start=1):
        has_visuals = image_count > 0 and index == 1
        pages.append({
            "page_number": index,
            "extracted_text": chunk.strip(),
            "has_visuals": has_visuals,
            "visual_area_ratio": 0.2 if has_visuals else 0.0,
            "content_quality": classify_page(chunk, has_visuals, 0.2 if has_visuals else 0.0),
            "_render": ("docx-embedded", index - 1),
        })
    return pages


def extract_docx_images(data, limit=3):
    import docx
    document = docx.Document(io.BytesIO(data))
    out = []
    for rel in document.part.rels.values():
        if "image" in rel.reltype:
            try:
                blob = rel.target_part.blob
                ext = (rel.target_part.partname.ext or "png").lstrip(".")
                out.append((blob, "image/%s" % ext))
            except Exception:
                continue
        if len(out) >= limit:
            break
    return out


# ------------------------------------------------------------------ Text / source
def _parse_text(data, ext):
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except (UnicodeDecodeError, LookupError):
            continue
    else:
        raise DocumentParseError("This file's text encoding could not be read.")

    if is_source_file(ext):
        # Source formatting is meaningful — keep whole and never send to Vision.
        lines = text.splitlines()
        pages = []
        per_page = 200
        for index, start in enumerate(range(0, max(len(lines), 1), per_page), start=1):
            body = "\n".join(lines[start:start + per_page])
            pages.append({
                "page_number": index,
                "extracted_text": body,
                "has_visuals": False,
                "visual_area_ratio": 0.0,
                "content_quality": QUALITY_TEXT_RICH if body.strip() else QUALITY_EMPTY,
                "is_source_code": True,
                "language": ext,
                "_render": None,
            })
        return pages

    pages = []
    size = 2500
    chunks = [text[i:i + size] for i in range(0, max(len(text), 1), size)] or [""]
    for index, chunk in enumerate(chunks, start=1):
        pages.append({
            "page_number": index,
            "extracted_text": chunk.strip(),
            "has_visuals": False,
            "visual_area_ratio": 0.0,
            "content_quality": classify_page(chunk, False),
            "_render": None,
        })
    return pages


# ------------------------------------------------------------------ Entry point
def parse(data, filename, declared_mime=None):
    """
    Returns (extension, pages). Raises UnsupportedFileError / DocumentParseError.
    """
    ext = detect_extension(filename, declared_mime)

    if ext == "pdf":
        pages = _parse_pdf(data)
    elif ext in ("pptx", "ppt"):
        if ext == "ppt":
            raise UnsupportedFileError(
                "Legacy .ppt files aren't supported. Save the deck as .pptx and upload again."
            )
        pages = _parse_pptx(data)
    elif ext in ("docx", "doc"):
        if ext == "doc":
            raise UnsupportedFileError(
                "Legacy .doc files aren't supported. Save the document as .docx and upload again."
            )
        pages = _parse_docx(data)
    else:
        pages = _parse_text(data, ext)

    pages = [p for p in pages if p["content_quality"] != QUALITY_EMPTY] or pages[:1]
    return ext, pages


def sniff_mime(data, ext):
    """Magic-number check — the filename alone is never trusted."""
    signatures = {
        "pdf": [b"%PDF-"],
        "pptx": [b"PK\x03\x04"],
        "docx": [b"PK\x03\x04"],
    }
    expected = signatures.get(ext)
    if not expected:
        return True
    head = data[:8]
    return any(head.startswith(sig) for sig in expected)
